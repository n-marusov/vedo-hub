"""Parser for DOCX documents using python-docx.

Extracts headings (Heading 1-6 styles), paragraphs, tables (→ datatype properties),
and document metadata.
"""

from __future__ import annotations

import logging
import os
from typing import Any

from parsers.base import BaseParser, ParseError
from parsers.models import ParsedDocument, Section, Table, TableCell, TableRow

logger = logging.getLogger("document-extractor.parsers.docx")

# Mapping from python-docx heading style names to levels
HEADING_STYLE_MAP: dict[str, int] = {
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Heading 5": 5,
    "Heading 6": 6,
}


class DocxParser(BaseParser):
    """Parser for DOCX files using python-docx."""

    SUPPORTED_EXTENSIONS = [".docx", ".docm"]

    async def parse(self, file_path: str) -> ParsedDocument:
        """Parse a DOCX file and extract structured content."""
        self.validate_extension(file_path)
        file_size = os.path.getsize(file_path)

        logger.debug(
            "Parsing DOCX: %s (size=%d bytes)",
            os.path.basename(file_path),
            file_size,
        )

        try:
            from docx import Document as DocxDocument
            from docx.opc.exceptions import PackageNotFoundError
        except ImportError as exc:
            raise ParseError(
                "python-docx not installed; run 'pip install python-docx'",
                original=exc,
            )

        try:
            doc = DocxDocument(file_path)
        except PackageNotFoundError:
            raise ParseError(f"Not a valid DOCX file: {file_path}")
        except Exception as exc:
            raise ParseError(f"Failed to open DOCX {file_path}: {exc}", original=exc)

        # ─── Process paragraphs in document order ───────────────────────────

        sections: list[Section] = []
        current_section = Section(heading="", level=0)
        current_paragraphs: list[str] = []
        current_tables: list[Table] = []

        for para in doc.paragraphs:
            style_name = para.style.name if para.style else ""
            text = para.text.strip()

            if not text:
                continue

            if style_name in HEADING_STYLE_MAP:
                # Flush previous section
                if current_paragraphs or current_tables:
                    sections.append(
                        Section(
                            heading=current_section.heading,
                            level=current_section.level,
                            paragraphs=current_paragraphs,
                            tables=current_tables,
                        )
                    )
                    current_paragraphs = []
                    current_tables = []

                current_section = Section(
                    heading=text,
                    level=HEADING_STYLE_MAP[style_name],
                )
            elif style_name.startswith("List") or style_name.startswith("List Bullet"):
                current_paragraphs.append(f"- {text}")
            elif style_name.startswith("List Number"):
                current_paragraphs.append(f"1. {text}")
            else:
                current_paragraphs.append(text)

        # ─── Process tables ─────────────────────────────────────────────────

        for table_idx, table in enumerate(doc.tables):
            parsed_table = Table(caption=f"table_{table_idx}")
            for row_idx, row in enumerate(table.rows):
                table_row = TableRow(is_header=(row_idx == 0))
                for cell in row.cells:
                    table_row.cells.append(TableCell(text=cell.text.strip()))
                parsed_table.rows.append(table_row)
            current_tables.append(parsed_table)

        # ─── Flush last section ─────────────────────────────────────────────

        if current_paragraphs or current_tables:
            sections.append(
                Section(
                    heading=current_section.heading,
                    level=current_section.level,
                    paragraphs=current_paragraphs,
                    tables=current_tables,
                )
            )

        # If still no sections, create one with all content
        if not sections:
            all_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            all_tables: list[Table] = []
            for table_idx, table in enumerate(doc.tables):
                parsed_table = Table(caption=f"table_{table_idx}")
                for row_idx, row in enumerate(table.rows):
                    table_row = TableRow(is_header=(row_idx == 0))
                    for cell in row.cells:
                        table_row.cells.append(TableCell(text=cell.text.strip()))
                    parsed_table.rows.append(table_row)
                all_tables.append(parsed_table)

            sections.append(
                Section(heading="", level=0, paragraphs=all_paragraphs, tables=all_tables)
            )

        # ─── Title & metadata ───────────────────────────────────────────────

        title = os.path.splitext(os.path.basename(file_path))[0]
        try:
            if doc.core_properties.title:
                title = doc.core_properties.title
        except Exception:
            pass

        metadata: dict[str, Any] = {"file_size_bytes": file_size, "section_count": len(sections)}
        try:
            if doc.core_properties.author:
                metadata["author"] = doc.core_properties.author
            if doc.core_properties.created:
                metadata["created"] = str(doc.core_properties.created)
            if doc.core_properties.modified:
                metadata["modified"] = str(doc.core_properties.modified)
        except Exception:
            pass

        logger.debug(
            "Parsed DOCX %s: %d sections, %d tables",
            os.path.basename(file_path),
            len(sections),
            len(current_tables),
        )

        return ParsedDocument(
            title=title,
            format="docx",
            sections=sections,
            metadata=metadata,
        )
